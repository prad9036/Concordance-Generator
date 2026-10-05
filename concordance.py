import re
import sys
import shutil
import zipfile
from collections import defaultdict, Counter
from pathlib import Path
import fitz
import pymupdf
import nltk
from nltk.stem import WordNetLemmatizer

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


# ============================================================
# SETTINGS
# ============================================================

MIN_WORD_LENGTH = 3

FONT_NAME = "Times New Roman"
FONT_SIZE = 9

LEMMATIZER = WordNetLemmatizer()
LEMMATIZATION_WARNING_SHOWN = False
NLTK_DATA_READY = False
NLTK_DATA_AVAILABLE = False

TOP_MARGIN = 0.55
BOTTOM_MARGIN = 0.55
LEFT_MARGIN = 0.60
RIGHT_MARGIN = 0.60

COLUMN_GAP = 0.25


def load_stop_words():
    stop_word_file = Path(__file__).with_name("stop_word_list.txt")
    return {
        line.split("#", 1)[0].strip().lower()
        for line in stop_word_file.read_text(encoding="utf-8").splitlines()
        if line.split("#", 1)[0].strip()
    }


STOP_WORDS = load_stop_words()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_pages(pdf_file):

    doc = fitz.open(pdf_file)

    pages = []

    print("\n[1/3] Reading PDF")

    print(f"Total PDF pages: {len(doc)}")

    for number, page in enumerate(doc, start=1):

        text = page.get_text("text")
        page_label = extract_printed_page_number(page)

        pages.append({
            "page": number,
            "page_label": page_label,
            "text": text
        })

        if number % 25 == 0:
            print(
                f"  Extracted {number}/{len(doc)} pages"
            )

    fill_missing_page_labels(pages)

    doc.close()

    return pages


def extract_printed_page_number(page):
    """Read a page number only when it is isolated in a page margin."""
    page_height = page.rect.height
    margin_blocks = []
    for block in page.get_text("blocks"):
        if len(block) <= 4:
            continue
        x0, y0, x1, y1, block_text = block[:5]
        if y0 <= page_height * 0.16 or y1 >= page_height * 0.84:
            margin_blocks.append((x0, y0, x1, y1, block_text.strip()))

    # Page numbers are isolated text blocks. Requiring a full block match
    # avoids interpreting DOI fragments, dates, and table values as folios.
    for x0, y0, x1, y1, block_text in margin_blocks:
        at_outer_edge = y0 <= page_height * 0.08 or y1 >= page_height * 0.92
        if at_outer_edge and y1 - y0 <= 24 and re.fullmatch(
            r"\d{1,4}(?:\.\d{1,3})?", block_text
        ):
            return block_text

    # Dotted running headers may share a line with the section title. Only
    # inspect the top margin and reject URLs or lines with other numeric data.
    for x0, y0, x1, y1, block_text in margin_blocks:
        if y0 > page_height * 0.16 or re.search(r"https?://|doi", block_text, re.IGNORECASE):
            continue
        matches = re.findall(r"(?<![\d.])(\d{1,3}\.\d{1,3})(?![\d.])", block_text)
        if len(matches) == 1:
            return matches[0]

    # Chapter opener pages in this PDF show "N CHAPTER" instead of N.1.
    for x0, y0, x1, y1, block_text in margin_blocks:
        match = re.search(r"\b(\d{1,3})\s+CHAPTER\b", block_text, re.IGNORECASE)
        if match:
            return f"{match.group(1)}.1"

    # Fall back to an explicit PDF page label if present.
    try:
        label = page.get_label()
        if label:
            return str(label)
    except (AttributeError, RuntimeError):
        pass

    return None


def fill_missing_page_labels(pages):
    """Infer omitted labels from adjacent sequential labels when possible."""
    known = []
    for index, item in enumerate(pages):
        label = item.get("page_label")
        if label and re.fullmatch(r"\d+(?:\.\d+)?", str(label)):
            parts = tuple(map(int, str(label).split(".")))
            known.append((index, parts))

    for (left_index, left_parts), (right_index, right_parts) in zip(known, known[1:]):
        distance = right_index - left_index
        if len(left_parts) == len(right_parts) == 1:
            if right_parts[0] - left_parts[0] != distance:
                continue
            for offset in range(1, distance):
                pages[left_index + offset]["page_label"] = str(left_parts[0] + offset)
        elif len(left_parts) == len(right_parts) == 2:
            major, minor = left_parts
            right_major, right_minor = right_parts
            if major != right_major or right_minor - minor != distance:
                continue
            for offset in range(1, distance):
                pages[left_index + offset]["page_label"] = f"{major}.{minor + offset}"

    # If the first visible folio follows an unnumbered first page, infer the
    # preceding sequence when that produces a positive page number.
    first_known = next(
        ((index, parts) for index, parts in known if not pages[index].get("page_label") is None),
        None,
    )
    if first_known:
        index, parts = first_known
        if len(parts) == 1 and parts[0] > index:
            for offset in range(index):
                pages[offset]["page_label"] = str(parts[0] - (index - offset))
        elif len(parts) == 2 and parts[1] > index:
            for offset in range(index):
                pages[offset]["page_label"] = f"{parts[0]}.{parts[1] - (index - offset)}"

    for item in pages:
        if not item.get("page_label"):
            item["page_label"] = str(item["page"])


# ============================================================
# DETECT REPEATED HEADERS / FOOTERS
# ============================================================

def find_repeated_lines(pages):

    counter = Counter()

    for item in pages:

        lines = item["text"].splitlines()

        # Look at lines near the top and bottom.
        candidates = (
            lines[:4] +
            lines[-4:]
        )

        for line in candidates:

            line = re.sub(
                r"\s+",
                " ",
                line
            ).strip().lower()

            if len(line) >= 4:

                counter[line] += 1


    # A line appearing on at least 10% of pages
    # is probably a running header/footer.

    threshold = max(
        3,
        int(len(pages) * 0.10)
    )

    repeated = {
        line
        for line, count in counter.items()
        if count >= threshold
    }

    return repeated


def remove_repeated_lines(
    text,
    repeated_lines
):

    result = []

    for line in text.splitlines():

        normalized = re.sub(
            r"\s+",
            " ",
            line
        ).strip().lower()

        if normalized in repeated_lines:
            continue

        result.append(line)

    return "\n".join(result)


# ============================================================
# CLEAN PDF TEXT
# ============================================================

def clean_text(text):

    # Join words split at the end of a line.
    #
    # naviga-
    # tion
    #
    # becomes:
    #
    # navigation

    text = re.sub(
        r"(\w)-\s*\n\s*(\w)",
        r"\1\2",
        text
    )


    # Convert line breaks to spaces.

    text = re.sub(
        r"\s*\n\s*",
        " ",
        text
    )


    # Collapse multiple spaces.

    text = re.sub(
        r"\s+",
        " ",
        text
    )


    return text.strip()


# ============================================================
# EXTRACT WORDS
# ============================================================

def extract_words(text):

    return re.findall(
        r"\b[A-Za-z]+(?:['’\-][A-Za-z]+)*\b",
        text
    )


def normalize_word(word):

    word = word.lower()

    # Normalize curly apostrophe.

    word = word.replace(
        "’",
        "'"
    )


    # Remove possessive.

    if word.endswith("'s"):

        word = word[:-2]


    # Remove leading/trailing punctuation.

    word = word.strip(
        "-'"
    )

    # Canonicalize this regular British/American spelling difference.
    if word.endswith("isation"):
        word = word[:-7] + "ization"

    return word


def ensure_nltk_data():
    """Make the required NLTK datasets available and report setup progress."""
    global NLTK_DATA_READY, NLTK_DATA_AVAILABLE
    if NLTK_DATA_READY:
        return NLTK_DATA_AVAILABLE

    data_dir = Path.cwd() / "nltk_data"
    data_dir.mkdir(parents=True, exist_ok=True)
    data_dir_string = str(data_dir)
    if data_dir_string not in nltk.data.path:
        nltk.data.path.insert(0, data_dir_string)

    print("\n[SETUP] Checking language resources")
    print(f"        Download location: {data_dir}")
    resources = (
        {
            "name": "wordnet",
            "package": "wordnet",
            "paths": ("corpora/wordnet", "corpora/wordnet.zip"),
            "files": (data_dir / "corpora" / "wordnet.zip",),
            "directories": (data_dir / "corpora" / "wordnet",),
        },
        {
            "name": "averaged_perceptron_tagger_eng",
            "package": "averaged_perceptron_tagger_eng",
            "paths": (
                "taggers/averaged_perceptron_tagger_eng",
                "taggers/averaged_perceptron_tagger_eng.zip",
            ),
            "files": (data_dir / "taggers" / "averaged_perceptron_tagger_eng.zip",),
            "directories": (
                data_dir / "taggers" / "averaged_perceptron_tagger_eng",
            ),
        },
    )
    total_resources = len(resources)
    all_ready = True
    for index, resource in enumerate(resources, start=1):
        progress = f"[NLTK {index}/{total_resources}]"
        if nltk_resource_is_valid(resource, data_dir_string):
            print(f"{progress} Ready: {resource['name']}")
            continue

        # A previous interrupted download can leave zero-byte or invalid ZIP
        # files. Remove only this resource's local files before retrying.
        remove_nltk_resource_files(resource)
        print(f"{progress} Downloading {resource['package']}...")
        downloaded = False
        try:
            downloaded = download_nltk_package(
                resource["package"], data_dir_string, progress
            )
        except Exception as error:
            print(f"{progress} Download failed: {error}")

        valid_after_download = downloaded and nltk_resource_is_valid(
            resource, data_dir_string
        )
        if valid_after_download:
            print(f"{progress} Installed: {resource['name']}")
        else:
            remove_nltk_resource_files(resource)
            print(f"{progress} Unavailable: {resource['name']}")
            all_ready = False

    NLTK_DATA_READY = True
    NLTK_DATA_AVAILABLE = all_ready
    return all_ready


def nltk_resource_is_valid(resource, data_dir):
    """Check that a resource exists locally and can actually be loaded."""
    try:
        found = False
        for resource_path in resource["paths"]:
            try:
                nltk.data.find(resource_path, paths=[data_dir])
                found = True
                break
            except LookupError:
                continue
        if not found:
            return False

        if resource["name"] == "wordnet":
            from nltk.corpus import wordnet
            wordnet.ensure_loaded()
            return wordnet.morphy("calculations", wordnet.NOUN) == "calculation"

        nltk.tag.PerceptronTagger(lang="eng")
        return True
    except (LookupError, OSError, ValueError, zipfile.BadZipFile):
        return False


def remove_nltk_resource_files(resource):
    """Remove corrupt or partial files for one managed NLTK package."""
    for file_path in resource["files"]:
        if file_path.exists():
            file_path.unlink()
    for directory in resource["directories"]:
        if directory.exists():
            shutil.rmtree(directory)


def download_nltk_package(package, download_dir, progress_label):
    """Download a package while showing its byte-transfer percentage."""
    downloader = nltk.downloader.Downloader(download_dir=download_dir)
    last_percent = None
    installed = False

    for message in downloader.incr_download(package, download_dir=download_dir):
        if isinstance(message, nltk.downloader.ProgressMessage):
            percent = max(0, min(100, int(message.progress)))
            if percent != last_percent:
                bar_width = 24
                filled = round(bar_width * percent / 100)
                bar = "#" * filled + "-" * (bar_width - filled)
                print(
                    f"\r{progress_label} {package}: "
                    f"[{bar}] {percent:3d}%",
                    end="",
                    flush=True,
                )
                last_percent = percent
        elif isinstance(message, nltk.downloader.ErrorMessage):
            if last_percent is not None:
                print()
            print(f"{progress_label} Download error: {message.message}")
        elif isinstance(message, nltk.downloader.FinishPackageMessage):
            installed = True

    if last_percent is not None:
        print()
    return installed


def lemmatize_words(words):
    """Lemmatize known English inflections; leave unknown words untouched."""
    global LEMMATIZATION_WARNING_SHOWN

    words = [normalize_word(word) for word in words]

    try:
        if not ensure_nltk_data():
            if not LEMMATIZATION_WARNING_SHOWN:
                print("NLTK resources are unavailable; words will remain unmodified.")
                LEMMATIZATION_WARNING_SHOWN = True
            return words

        tagged_words = nltk.pos_tag(words)
        lemmas = []

        for word, tag in tagged_words:
            pos = tag[0]
            wordnet_pos = {
                "J": "a",
                "N": "n",
                "R": "r",
                "V": "v",
            }.get(pos)

            if wordnet_pos is None:
                lemma = word
            else:
                lemma = LEMMATIZER.lemmatize(word, wordnet_pos)

            # POS tagging can be unreliable in extracted technical text.
            # Try the likely grammatical category for common inflectional
            # endings, accepting a change only when WordNet recognizes a lemma.
            fallback_pos = []
            if word.endswith(("s", "es")):
                fallback_pos.extend(("n", "v"))
            if word.endswith(("ing", "ed", "en")):
                fallback_pos.append("v")
            if word.endswith(("er", "est")):
                fallback_pos.append("a")

            if lemma == word:
                for candidate_pos in fallback_pos:
                    candidate = LEMMATIZER.lemmatize(word, candidate_pos)
                    if candidate != word:
                        lemma = candidate
                        break

            lemmas.append(lemma)

        return lemmas
    except (LookupError, OSError, zipfile.BadZipFile) as error:
        if not LEMMATIZATION_WARNING_SHOWN:
            print(
                "NLTK lemmatization data is unavailable; words will remain "
                f"unmodified. Details: {error}"
            )
            LEMMATIZATION_WARNING_SHOWN = True
        return words


# ============================================================
# BUILD CONCORDANCE
# ============================================================

def build_concordance(pages):

    print()
    print("\n[2/3] Building concordance")


    repeated_lines = (
        find_repeated_lines(pages)
    )


    print(
        f"Detected {len(repeated_lines)} "
        "repeated header/footer lines."
    )


    # word -> set of page numbers

    concordance = defaultdict(set)

    total_words = 0


    for index, item in enumerate(
        pages,
        start=1
    ):

        page_number = item.get("page_label", item["page"])


        text = remove_repeated_lines(
            item["text"],
            repeated_lines
        )


        text = clean_text(
            text
        )


        if not text:
            continue


        words = extract_words(
            text
        )

        words = [normalize_word(word) for word in words]
        words = lemmatize_words(words)


        for raw_word in words:

            word = normalize_word(
                raw_word
            )


            if len(word) < MIN_WORD_LENGTH:
                continue


            if word in STOP_WORDS:
                continue


            concordance[word].add(
                page_number
            )


            total_words += 1


        if index % 25 == 0:

            print(
                f"  Indexed "
                f"{index}/{len(pages)} pages"
            )


    return (
        concordance,
        total_words
    )


# ============================================================
# FORMAT PAGE NUMBERS
# ============================================================

def format_page_numbers(pages):

    if not pages:
        return ""

    def page_sort_key(value):
        value = str(value)
        if value.isdigit():
            return (0, int(value), 0, "")
        dotted = re.fullmatch(r"(\d+)\.(\d+)", value)
        if dotted:
            return (1, int(dotted.group(1)), int(dotted.group(2)), "")
        return (2, 0, 0, value)

    return ", ".join(sorted({str(value) for value in pages}, key=page_sort_key))


# ============================================================
# SET WORD COLUMN LAYOUT
# ============================================================

def set_three_columns(section):

    sectPr = section._sectPr


    # Find existing <w:cols>

    cols_list = sectPr.xpath(
        "./w:cols"
    )


    if cols_list:

        cols = cols_list[0]

    else:

        cols = OxmlElement(
            "w:cols"
        )

        sectPr.append(
            cols
        )


    # Three columns.

    cols.set(
        qn("w:num"),
        "3"
    )


    # Equal width.

    cols.set(
        qn("w:equalWidth"),
        "1"
    )


    # Space between columns.
    #
    # Word uses twips here.
    # 360 twips = 0.25 inch.

    cols.set(
        qn("w:space"),
        "360"
    )


# ============================================================
# CREATE WORD DOCUMENT
# ============================================================

def create_docx(
    concordance,
    output_file,
    pdf_name
):

    print()
    print("\n[3/3] Creating Word document")


    doc = Document()


    # ========================================================
    # PAGE SETUP
    # ========================================================

    section = doc.sections[0]


    section.top_margin = Inches(
        TOP_MARGIN
    )

    section.bottom_margin = Inches(
        BOTTOM_MARGIN
    )

    section.left_margin = Inches(
        LEFT_MARGIN
    )

    section.right_margin = Inches(
        RIGHT_MARGIN
    )


    # Native Word columns.

    set_three_columns(
        section
    )


    # ========================================================
    # DEFAULT FONT
    # ========================================================

    normal = doc.styles["Normal"]

    normal.font.name = FONT_NAME

    normal.font.size = Pt(
        FONT_SIZE
    )

    normal._element.rPr.rFonts.set(
        qn("w:eastAsia"),
        FONT_NAME
    )


    # ========================================================
    # TITLE
    # ========================================================

    title = doc.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    title.paragraph_format.space_after = Pt(2)


    run = title.add_run(
        "CONCORDANCE"
    )

    run.bold = True

    run.font.name = FONT_NAME

    run.font.size = Pt(15)


    # Subtitle

    subtitle = doc.add_paragraph()

    subtitle.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    subtitle.paragraph_format.space_after = Pt(7)


    run = subtitle.add_run(
        pdf_name
    )

    run.italic = True

    run.font.name = FONT_NAME

    run.font.size = Pt(8)


    # ========================================================
    # CONCORDANCE
    # ========================================================

    words = sorted(
        concordance.keys()
    )


    current_letter = None


    for word in words:

        first_letter = (
            word[0].upper()
        )


        # ----------------------------------------------------
        # LETTER HEADING
        # ----------------------------------------------------

        if first_letter != current_letter:

            current_letter = first_letter


            p = doc.add_paragraph()


            p.paragraph_format.space_before = Pt(5)

            p.paragraph_format.space_after = Pt(2)

            p.paragraph_format.keep_with_next = True

            p.paragraph_format.keep_together = True


            run = p.add_run(
                current_letter
            )

            run.bold = True

            run.font.name = FONT_NAME

            run.font.size = Pt(11)


        # ----------------------------------------------------
        # WORD ENTRY
        # ----------------------------------------------------

        p = doc.add_paragraph()


        p.paragraph_format.space_before = Pt(0)

        p.paragraph_format.space_after = Pt(1)

        p.paragraph_format.line_spacing = 1.0

        p.paragraph_format.keep_together = True


        # ----------------------------------------------------
        # WORD
        # ----------------------------------------------------

        run = p.add_run(
            word
        )

        run.bold = True

        run.font.name = FONT_NAME

        run.font.size = Pt(
            FONT_SIZE
        )


        # ----------------------------------------------------
        # PAGE NUMBERS
        # ----------------------------------------------------

        p.add_run(
            "  "
        )


        page_numbers = (
            format_page_numbers(
                concordance[word]
            )
        )


        run = p.add_run(
            page_numbers
        )

        run.font.name = FONT_NAME

        run.font.size = Pt(
            FONT_SIZE
        )


    # ========================================================
    # FOOTER
    # ========================================================

    footer = section.footer


    footer_p = footer.paragraphs[0]


    footer_p.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    run = footer_p.add_run(
        "Concordance"
    )


    run.font.name = FONT_NAME

    run.font.size = Pt(8)


    # ========================================================
    # SAVE
    # ========================================================

    doc.save(
        output_file
    )


    print()
    print(
        f"Saved: {output_file}"
    )


# ============================================================
# MAIN
# ============================================================

def main():

    if len(sys.argv) not in (2, 3):

        print()
        print(
            "Usage:"
        )

        print(
            'python concordance.py '
            '"book.pdf" ["output.docx"]'
        )

        print()

        sys.exit(1)


    pdf_file = sys.argv[1]
    if len(sys.argv) == 3:
        output_file = sys.argv[2]
    else:
        input_path = Path(pdf_file)
        output_file = str(
            Path.cwd() / f"concordance_{input_path.stem}.docx"
        )


    print()
    print("\n" + "=" * 56)
    print("  BOOK CONCORDANCE GENERATOR")
    print("=" * 56)
    print(f"  Source: {pdf_file}")
    print(f"  Output: {output_file}")

    ensure_nltk_data()


    # --------------------------------------------------------
    # READ PDF
    # --------------------------------------------------------

    pages = extract_pages(
        pdf_file
    )


    # --------------------------------------------------------
    # BUILD INDEX
    # --------------------------------------------------------

    concordance, total_words = (
        build_concordance(
            pages
        )
    )


    # --------------------------------------------------------
    # CREATE WORD FILE
    # --------------------------------------------------------

    create_docx(
        concordance,
        output_file,
        pdf_file
    )


    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    print()
    print("=" * 60)
    print("DONE")
    print("=" * 60)

    print(
        f"Unique indexed words: "
        f"{len(concordance):,}"
    )

    print(
        f"Word occurrences processed: "
        f"{total_words:,}"
    )

    print(
        f"Output: {output_file}"
    )

    print()


# ============================================================
# RUN
# ============================================================

if __name__ == "__main__":
    main()
